# -*- coding: utf-8 -*-
"""只看 WORLDBOOK-NAME-INDEX.xlsx 的结构：sheet 名 / 表头 / 行数 / 前几行样例。"""
import io
import openpyxl

F = r"D:/AWAKE-Dev/AWAKE/docs/worldbook-migration/WORLDBOOK-NAME-INDEX.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)
print("sheet 列表:", wb.sheetnames)
for ws in wb.worksheets:
    print()
    print("=" * 70)
    print("sheet: %s  维度: %s  行数: %d  列数: %d" % (ws.title, ws.dimensions, ws.max_row, ws.max_column))
    print("=" * 70)
    for r in ws.iter_rows(min_row=1, max_row=min(6, ws.max_row), values_only=True):
        print("   " + " | ".join(("" if c is None else str(c))[:26] for c in r))
    print("   ...")
    for r in ws.iter_rows(min_row=max(1, ws.max_row - 2), max_row=ws.max_row, values_only=True):
        print("   " + " | ".join(("" if c is None else str(c))[:26] for c in r))
