# -*- coding: utf-8 -*-
"""临时：核验两层版工作簿结构（名录列序/筛选/格式 + 分类总览 + 分类定义）。"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
wb = openpyxl.load_workbook(XLSX)
print("sheets:", wb.sheetnames)

ws = wb["名录"]
print(f"\n名录 dims={ws.dimensions} freeze={ws.freeze_panes} filter={ws.auto_filter.ref}")
print("header:", [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)])
for r in (2, 3):
    print(f"  r{r}:", [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])
print("条件格式:", [(str(k.sqref), v[0].operator, v[0].formula)
                    for k, v in ws.conditional_formatting._cf_rules.items()])

for name in ("分类总览", "分类定义", "域汇总"):
    w = wb[name]
    print(f"\n=== {name} dims={w.dimensions} ===")
    for r in range(1, w.max_row + 1):
        vals = [w.cell(row=r, column=c).value for c in range(1, w.max_column + 1)]
        vals = ["" if v is None else str(v).replace("\n", " / ") for v in vals]
        line = " | ".join(vals).rstrip(" |")
        if line.strip():
            print(f"{r:>3} | {line[:160]}")
