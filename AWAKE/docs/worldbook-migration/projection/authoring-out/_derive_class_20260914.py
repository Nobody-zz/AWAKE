# -*- coding: utf-8 -*-
"""临时：从 yaml 派生 类别/子类，并与 xlsx「名录」逐行比对，锁定派生规则。"""
import io, os, glob, collections
import yaml
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))

# 规则A：doc_id 第三段整体
# 规则B：doc_id 第三段首个 '-' 之前 token
a = collections.Counter(); b = collections.Counter()
rows = {}
for path in sorted(glob.glob(os.path.join(HERE, "*.yaml"))):
    d = yaml.safe_load(io.open(path, encoding="utf-8"))
    if not str(d.get("id", "")).startswith("doc."):
        continue
    seg = d["id"].split(".", 2)[2]
    a[seg] += 1
    b[seg.split("-", 1)[0]] += 1
    rows[d["id"]] = (d.get("domain", ""), seg)

print("--- A: 第三段整体 (top 40) ---")
for k, v in a.most_common(40):
    print(f"  {k:40s} {v}")
print(f"  distinct={len(a)} total={sum(a.values())}")

print("\n--- B: 首个 '-' 前 token ---")
for k, v in sorted(b.items()):
    print(f"  {k:20s} {v}")
print(f"  distinct={len(b)} total={sum(b.values())}")

# 与 xlsx 比对
wb = openpyxl.load_workbook(XLSX)
ws = wb["名录"]
mism = []
for r in range(2, ws.max_row + 1):
    docid = ws.cell(row=r, column=4).value
    cat = ws.cell(row=r, column=2).value
    sub = ws.cell(row=r, column=3).value
    if docid not in rows:
        mism.append((docid, "NOT-IN-YAML")); continue
    dom, seg = rows[docid]
    exp_sub = seg.split("-", 1)[0]
    if exp_sub == "mount" or exp_sub == "mountains":
        exp_sub = "mountain"
    if cat != dom or sub != exp_sub:
        mism.append((docid, f"xlsx=({cat},{sub}) expect=({dom},{exp_sub})"))
print(f"\n--- xlsx 比对 ---  rows={ws.max_row-1}  mismatches={len(mism)}")
for m in mism[:30]:
    print("  ", m)
