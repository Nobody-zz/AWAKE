# -*- coding: utf-8 -*-
"""临时：校验新 xlsx 的条件格式/百分比/列宽落点。"""
import os
import openpyxl
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
NEW = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
wb = openpyxl.load_workbook(NEW)

ws = wb["名录"]
print("名录 条件格式:")
for rng, rules in ws.conditional_formatting._cf_rules.items():
    for r in rules:
        print("  ", rng.sqref, "->", r.operator, r.formula)
print("名录 列宽:", {get_column_letter(c): ws.column_dimensions[get_column_letter(c)].width
                    for c in range(1, ws.max_column + 1)})
print("名录 J2 =", ws.cell(row=2, column=10).value, " R2 =", ws.cell(row=2, column=18).value)

t = wb["分类总览"]
print("\n分类总览 百分比列 number_format:")
for r in range(1, t.max_row + 1):
    v = t.cell(row=r, column=4).value
    if isinstance(v, float):
        print(f"  r{r}: {v!r} fmt={t.cell(row=r, column=4).number_format!r}")

w = wb["刷新说明"]
print("\n刷新说明 行数:", w.max_row)
for r in range(1, w.max_row + 1):
    print("  ", r, str(w.cell(row=r, column=1).value)[:60])
