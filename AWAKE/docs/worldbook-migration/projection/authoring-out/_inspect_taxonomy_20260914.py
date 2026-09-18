# -*- coding: utf-8 -*-
"""临时：导出「分类总览」表全部单元格（只读）。"""
import os
import openpyxl
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
wb = openpyxl.load_workbook(XLSX)
ws = wb["分类总览"]
for r in range(1, ws.max_row + 1):
    vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        vals.append("" if v is None else str(v))
    print(f"{r:>2} | " + " ‖ ".join(vals))
print("\n-- merged:", ws.merged_cells.ranges)
