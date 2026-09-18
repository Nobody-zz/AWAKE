# -*- coding: utf-8 -*-
"""临时：核验新工作簿（6 表）+ 名录未回归 + 导出分类定义表。"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
NEW = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
OLD = os.path.abspath(os.path.join(HERE, "..", "..", "_archive-WORLDBOOK-NAME-INDEX-20260914-sheetagent.xlsx"))

b = openpyxl.load_workbook(NEW)
a = openpyxl.load_workbook(OLD)
print("NEW sheets:", b.sheetnames)
print("OLD sheets:", a.sheetnames)

wa, wb_ = a["名录"], b["名录"]
diff = sum(1 for r in range(1, 450) for c in range(1, 19)
           if wa.cell(row=r, column=c).value != wb_.cell(row=r, column=c).value)
print(f"名录 448x18 差异格 = {diff}")
print("名录 filter:", wb_.auto_filter.ref, "| freeze:", wb_.freeze_panes)

w = b["分类定义"]
print(f"\n=== 分类定义 dims={w.dimensions} freeze={w.freeze_panes} ===")
for r in range(1, w.max_row + 1):
    vals = [w.cell(row=r, column=c).value for c in range(1, 8)]
    vals = ["" if v is None else str(v) for v in vals]
    line = " ‖ ".join(vals).rstrip(" ‖")
    if line.strip():
        print(f"{r:>3} | {line[:150]}")

ws5 = b["刷新说明"]
print("\n=== 刷新说明 ===")
for r in range(1, ws5.max_row + 1):
    v = ws5.cell(row=r, column=1).value
    if v:
        print(f"{r:>3} | {str(v)[:130]}")
