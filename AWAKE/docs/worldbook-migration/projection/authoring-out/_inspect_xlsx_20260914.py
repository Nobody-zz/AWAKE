# -*- coding: utf-8 -*-
"""临时：核对当前 WORLDBOOK-NAME-INDEX.xlsx 的表/列结构（只读）。"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
wb = openpyxl.load_workbook(XLSX)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("  freeze:", ws.freeze_panes, " autofilter:", ws.auto_filter.ref)
    hdr = [ws.cell(row=1, column=c).value for c in range(1, min(ws.max_column, 20) + 1)]
    print("  header:", hdr)
    # 列宽
    widths = {}
    for c in range(1, ws.max_column + 1):
        L = openpyxl.utils.get_column_letter(c)
        w = ws.column_dimensions[L].width
        if w:
            widths[L] = round(w, 1)
    print("  widths:", widths)
    if name == "名录" and ws.max_row > 1:
        for r in (2, 3):
            print(f"  row{r}:", [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 20) + 1)])
