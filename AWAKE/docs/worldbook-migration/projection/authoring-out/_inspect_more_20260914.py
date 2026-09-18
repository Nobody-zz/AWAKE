# -*- coding: utf-8 -*-
"""临时：导出 域汇总 / 刷新说明 内容（只读）。"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
wb = openpyxl.load_workbook(XLSX)
for name in ("域汇总", "刷新说明"):
    ws = wb[name]
    print(f"\n=== {name} ===")
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        vals = ["" if v is None else str(v) for v in vals]
        print(f"{r:>2} | " + " ‖ ".join(vals))
