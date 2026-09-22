# -*- coding: utf-8 -*-
"""读 NAME-INDEX 的「刷新说明」与「交叉覆盖」两个 sheet 全文 + 分类定义底部注脚。"""
import openpyxl

F = r"D:/AWAKE-Dev/AWAKE/docs/worldbook-migration/WORLDBOOK-NAME-INDEX.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)

for t in ("刷新说明", "交叉覆盖"):
    ws = wb[t]
    print("=" * 72)
    print("sheet:", t, ws.dimensions)
    print("=" * 72)
    for r in ws.iter_rows(values_only=True):
        line = " | ".join("" if c is None else str(c) for c in r)
        if line.strip(" |"):
            print(line)

ws = wb["分类定义"]
print()
print("=" * 72)
print("分类定义 · 尾部 6 行")
print("=" * 72)
rows = list(ws.iter_rows(values_only=True))
for r in rows[-6:]:
    print(" | ".join("" if c is None else str(c) for c in r))
