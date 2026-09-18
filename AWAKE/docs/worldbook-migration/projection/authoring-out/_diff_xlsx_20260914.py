# -*- coding: utf-8 -*-
"""临时：比对 重跑产物 vs sheet-agent 原版（名录全格 + 分类总览 文本）。"""
import os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
NEW = os.path.abspath(os.path.join(HERE, "..", "..", "WORLDBOOK-NAME-INDEX.xlsx"))
OLD = os.path.abspath(os.path.join(HERE, "..", "..", "_archive-WORLDBOOK-NAME-INDEX-20260914-sheetagent.xlsx"))

a = openpyxl.load_workbook(OLD)
b = openpyxl.load_workbook(NEW)
print("sheets OLD:", a.sheetnames)
print("sheets NEW:", b.sheetnames)
assert a.sheetnames == b.sheetnames, "sheet 列表不同"

# 名录全格比对
wa, wb_ = a["名录"], b["名录"]
print(f"名录 dims OLD={wa.dimensions} NEW={wb_.dimensions}")
diff = 0
for r in range(1, max(wa.max_row, wb_.max_row) + 1):
    for c in range(1, max(wa.max_column, wb_.max_column) + 1):
        va = wa.cell(row=r, column=c).value
        vb = wb_.cell(row=r, column=c).value
        if va != vb:
            diff += 1
            if diff <= 20:
                print(f"  名录 r{r}c{c}: OLD={va!r} NEW={vb!r}")
print(f"名录 差异格数 = {diff}")

# 冻结/筛选
print("名录 freeze OLD/NEW:", wa.freeze_panes, wb_.freeze_panes)
print("名录 filter OLD/NEW:", wa.auto_filter.ref, wb_.auto_filter.ref)

# 分类总览文本比对
ta, tb = a["分类总览"], b["分类总览"]
print(f"\n分类总览 dims OLD={ta.dimensions} NEW={tb.dimensions}")
tdiff = 0
for r in range(1, max(ta.max_row, tb.max_row) + 1):
    for c in range(1, max(ta.max_column, tb.max_column) + 1):
        va = ta.cell(row=r, column=c).value
        vb = tb.cell(row=r, column=c).value
        if isinstance(va, float) and isinstance(vb, float):
            same = abs(va - vb) < 1e-9
        else:
            same = va == vb
        if not same:
            tdiff += 1
            if tdiff <= 30:
                print(f"  总览 r{r}c{c}: OLD={va!r} NEW={vb!r}")
print(f"分类总览 差异格数 = {tdiff}")

# 域汇总 / 交叉覆盖 数字比对
for nm in ("域汇总", "交叉覆盖"):
    xa, xb = a[nm], b[nm]
    d = 0
    for r in range(1, max(xa.max_row, xb.max_row) + 1):
        for c in range(1, max(xa.max_column, xb.max_column) + 1):
            if xa.cell(row=r, column=c).value != xb.cell(row=r, column=c).value:
                d += 1
                if d <= 10:
                    print(f"  {nm} r{r}c{c}: OLD={xa.cell(row=r,column=c).value!r} NEW={xb.cell(row=r,column=c).value!r}")
    print(f"{nm} 差异格数 = {d}")
