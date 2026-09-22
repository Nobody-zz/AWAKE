# -*- coding: utf-8 -*-
"""核验刷新后的 NAME-INDEX：新档在不在名录、域汇总是否自洽、生成时间。"""
import openpyxl

F = r"D:/AWAKE-Dev/AWAKE/docs/worldbook-migration/WORLDBOOK-NAME-INDEX.xlsx"
wb = openpyxl.load_workbook(F, data_only=True)

ws = wb["名录"]
rows = list(ws.iter_rows(values_only=True))
print("名录：%d 行（表头 1 + 档案 %d）" % (len(rows), len(rows) - 1))
hdr = rows[0]
print("表头：", " | ".join(str(c) for c in hdr))

hit = [r for r in rows[1:] if r[4] == "doc.war.troops-cataphract"]
print()
print("--- 新档行 ---")
for r in hit:
    for h, c in zip(hdr, r):
        print("   %-12s %s" % (h, c))

print()
print("--- 域汇总 ---")
for r in wb["域汇总"].iter_rows(values_only=True):
    if r[0] in ("war", "合计", "主分类"):
        print("   " + " | ".join("" if c is None else str(c) for c in r))

print()
print("--- 名录尾部 3 行（看排序是否含新档）---")
for r in rows[-3:]:
    print("   %s | %s | %s" % (r[0], r[4], r[6]))

print()
print("--- 刷新说明末行 ---")
sn = list(wb["刷新说明"].iter_rows(values_only=True))
print("   " + str(sn[-1][0]))
